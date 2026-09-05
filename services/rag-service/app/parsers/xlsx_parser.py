"""Excel 解析器：openpyxl 按行切块（V5，design D2）。"""
import asyncio
from pathlib import Path

from app.parsers.base import DocumentParser
from common.config.settings import settings


class XlsxParser(DocumentParser):
    """Excel 解析：全部 sheet（KB_EXCEL_SHEET_ALL）/ 表头行（KB_EXCEL_HEADER_ROW）/ 行切块序列化。

    - 每行序列化为 "表格(第0页)[sheet名]: 列名1:值1, 列名2:值2,..."（sheet 名作定位前缀）；
    - 表头行之前的行视为标题跳过；空行与全空值行跳过；空单元格不序列化；
    - 损坏/加密 → 抛通用异常，由 ingest 沿既有失败路径转 A3001（或降级占位）；
    - 解析耗时受 KB_PARSE_TIMEOUT 约束（ingest 层 asyncio.wait_for 包裹）。
    """

    def _extract(self, file_path: Path) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError as e:  # pragma: no cover
            raise RuntimeError(f"openpyxl 未安装: {e}") from e

        try:
            # 双读取：data_only=True 取公式缓存计算值；data_only=False 取公式字符串/原始值，
            # 用于在公式单元格无缓存（读到 None）时回退读取，避免"有数据却被判空"（A3002）。
            # 注意：必须 read_only=False。read_only 模式依赖 sheet 的 <dimension> 声明，
            # 某些工具导出的 xlsx 该声明不完整（如只声明 A1），会误判为 1 行 1 列、丢光数据
            # （实测 200 条数据被读成 1 格，text_len=0 -> A3002）；非只读模式按 XML 完整解析不受影响。
            wb_values = load_workbook(str(file_path), read_only=False, data_only=True)
            wb_formula = load_workbook(str(file_path), read_only=False, data_only=False)
        except Exception as e:
            raise ValueError(f"Excel 文件损坏或加密: {e}") from e

        sheet_names = list(wb_values.sheetnames)
        if not sheet_names:
            wb_values.close()
            wb_formula.close()
            raise ValueError("Excel 无有效 sheet")
        if not settings.KB_EXCEL_SHEET_ALL:
            sheet_names = sheet_names[:1]

        header_row = max(1, settings.KB_EXCEL_HEADER_ROW)
        lines: list[str] = []
        try:
            for name in sheet_names:
                ws_values = wb_values[name]
                ws_formula = wb_formula[name]
                header: list[str] = []
                data_rows: list[tuple] = []
                for idx, (row_values, row_formula) in enumerate(
                    zip(
                        ws_values.iter_rows(values_only=True),
                        ws_formula.iter_rows(values_only=True),
                    ),
                    start=1,
                ):
                    if idx < header_row:
                        continue  # 表头之前的行视为标题跳过
                    # 回退取值：data_only=True 读到 None 时，用 data_only=False 的公式/原始值作可入库文本
                    row = [
                        (val if val is not None else fb)
                        for val, fb in zip(row_values, row_formula)
                    ]
                    if idx == header_row:
                        header = [
                            str(c).strip() if c is not None else f"列{i + 1}"
                            for i, c in enumerate(row)
                        ]
                        continue
                    data_rows.append(tuple(row))
                for row in data_rows:
                    if all(c is None or str(c).strip() == "" for c in row):
                        continue  # 空行跳过
                    pairs = [
                        f"{header[i]}:{str(c).strip()}"
                        for i, c in enumerate(row)
                        if i < len(header) and c is not None and str(c).strip() != ""
                    ]
                    if pairs:
                        lines.append(f"表格(第0页)[{name}]: {', '.join(pairs)}")
        finally:
            wb_values.close()
            wb_formula.close()
        return "\n".join(lines)

    async def extract_text(self, file_path: Path) -> str:
        return await asyncio.to_thread(self._extract, file_path)
