"""PPT 解析器：python-pptx 提取标题+正文+演讲者备注（V5，design D3）。"""
import asyncio
import re
from pathlib import Path

from app.parsers.base import DocumentParser

# 版式占位文案（精确匹配才剔除，保守启发式，防误删正文）
_PLACEHOLDER_TEXTS = frozenset({
    "点击添加标题", "单击此处添加标题", "点击此处添加标题", "单击以添加文本",
    "点击添加副标题", "单击此处添加副标题", "在此处输入文本", "点击添加文本",
    "Click to add title", "Click to add subtitle", "Click to add text",
    "Add title", "Add subtitle", "Add text", "Add an outline",
})
# 页码/页脚行（确定性模式才剔除）
_PAGE_NUMBER_RE = re.compile(r"^\s*\d+\s*$")
_SLASH_TOTAL_RE = re.compile(r"^\s*\d+\s*/\s*总页数\s*$")


class PptxParser(DocumentParser):
    """PPT 解析：逐页抽取标题(占位符)+正文(非标题文本框)+演讲者备注（并入页文本）。

    - 页码行（纯数字 / "N/总页数"）与精确匹配的版式占位文案保守过滤，正文含数字/短句不删；
    - 多页按序拼接；损坏/加密 → 抛通用异常，由 ingest 沿既有失败路径转 A3001。
    """

    def _extract(self, file_path: Path) -> str:
        try:
            from pptx import Presentation
        except ImportError as e:  # pragma: no cover
            raise RuntimeError(f"python-pptx 未安装: {e}") from e

        try:
            prs = Presentation(str(file_path))
        except Exception as e:
            raise ValueError(f"PPT 文件损坏或加密: {e}") from e

        pages: list[str] = []
        for i, slide in enumerate(prs.slides, start=1):
            title_shape = slide.shapes.title
            title_text = ""
            body_lines: list[str] = []
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = self._filter_text(shape.text_frame.text)
                if not text:
                    continue
                filtered_lines = [
                    ln
                    for ln in text.split("\n")
                    if not _PAGE_NUMBER_RE.match(ln.strip())
                    and not _SLASH_TOTAL_RE.match(ln.strip())
                ]
                joined = "\n".join(filtered_lines).strip()
                if not joined:
                    continue
                if title_shape is not None and shape == title_shape:
                    title_text = joined
                else:
                    body_lines.append(joined)
            # 演讲者备注并入页文本（缺失/异常不影响正文）
            try:
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        body_lines.append(f"备注: {notes}")
            except Exception:
                pass

            parts = [f"第{i}页"]
            if title_text:
                parts.append(f"标题: {title_text}")
            if body_lines:
                parts.append("\n".join(body_lines))
            pages.append("\n".join(parts))
        return "\n\n".join(pages)

    def _filter_text(self, text: str) -> str:
        """剔除与版式占位文案精确匹配的行（保守启发式）。"""
        lines = []
        for ln in text.split("\n"):
            stripped = ln.strip()
            if stripped and stripped not in _PLACEHOLDER_TEXTS:
                lines.append(stripped)
        return "\n".join(lines)

    async def extract_text(self, file_path: Path) -> str:
        return await asyncio.to_thread(self._extract, file_path)

    async def extract_with_assets(self, file_path: Path) -> tuple[str, list]:
        """V5：正文文本 + 内嵌图片/表格资产（design D4/D5）；资产失败仅丢弃不影响正文。"""
        from app.parsers.asset_extractor import AssetExtractor, assets_enabled

        text = await self.extract_text(file_path)
        if not assets_enabled():
            return text, []
        try:
            assets = await AssetExtractor().extract(file_path)
        except Exception:
            assets = []
        return text, assets
