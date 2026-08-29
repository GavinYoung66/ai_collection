"""PPT 解析器单元测试（V5，design D3；tasks 11.2）。

覆盖：标题+正文+备注并入页文本 / 页码行与版式占位过滤（不产生独立 chunk）/
正文含数字不误删 / 备注可检索命中。
"""
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.parsers.pptx_parser import PptxParser


def _make_pptx(path: Path, slides: list[dict]) -> None:
    """生成真实 pptx：每页 {title?, body?, notes?, body_lines?}。

    使用 layout 1（带标题占位符）；正文用文本框。
    """
    prs = Presentation()
    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if s.get("title"):
            slide.shapes.title.text = s["title"]
        if s.get("body"):
            tb = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
            tb.text = s["body"]
        if s.get("body_lines"):
            tb = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(2))
            tf = tb.text_frame
            tf.text = s["body_lines"][0]
            for ln in s["body_lines"][1:]:
                tf.add_paragraph().text = ln
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(str(path))


class TestPptxParser:
    @pytest.mark.asyncio
    async def test_title_body_notes_merged_into_page_text(self, tmp_path):
        """标题+正文+演讲者备注并入页文本。"""
        f = tmp_path / "basic.pptx"
        _make_pptx(f, [{
            "title": "电话催收合规",
            "body": "礼貌开场，不得辱骂",
            "notes": "开场需报工号",
        }])
        text = await PptxParser().extract_text(f)
        assert "第1页" in text
        assert "标题: 电话催收合规" in text
        assert "礼貌开场，不得辱骂" in text
        assert "备注: 开场需报工号" in text

    @pytest.mark.asyncio
    async def test_page_number_line_filtered(self, tmp_path):
        """纯数字页码行过滤，不产生独立 chunk。"""
        f = tmp_path / "pageno.pptx"
        _make_pptx(f, [{"body_lines": ["12", "正文内容"]}])
        text = await PptxParser().extract_text(f)
        assert "正文内容" in text
        # 纯数字行被过滤（不出现独立 "12" 行）
        assert "\n12\n" not in text and not text.endswith("12")

    @pytest.mark.asyncio
    async def test_slash_total_page_line_filtered(self, tmp_path):
        """"N/总页数" 页码行过滤。"""
        f = tmp_path / "slash.pptx"
        _make_pptx(f, [{"body_lines": ["3/总页数", "内容"]}])
        text = await PptxParser().extract_text(f)
        assert "内容" in text
        assert "3/总页数" not in text

    @pytest.mark.asyncio
    async def test_placeholder_text_filtered(self, tmp_path):
        """版式占位文案（精确匹配）过滤，不产生独立 chunk。"""
        f = tmp_path / "placeholder.pptx"
        _make_pptx(f, [{"body_lines": ["单击此处添加标题", "真实正文"]}])
        text = await PptxParser().extract_text(f)
        assert "真实正文" in text
        assert "单击此处添加标题" not in text

    @pytest.mark.asyncio
    async def test_body_with_numbers_kept(self, tmp_path):
        """正文含数字不误删（"共 3 期，金额 5000 元"）。"""
        f = tmp_path / "numbers.pptx"
        _make_pptx(f, [{"body": "共 3 期，金额 5000 元，分 12 个月"}])
        text = await PptxParser().extract_text(f)
        assert "共 3 期，金额 5000 元，分 12 个月" in text

    @pytest.mark.asyncio
    async def test_notes_searchable(self, tmp_path):
        """备注并入页文本后可检索命中（source=ppt 场景的可检索载体）。"""
        f = tmp_path / "notes.pptx"
        _make_pptx(f, [{
            "body": "正常正文",
            "notes": "协商分期需登记在案",
        }])
        text = await PptxParser().extract_text(f)
        assert "备注: 协商分期需登记在案" in text

    @pytest.mark.asyncio
    async def test_multiple_pages_in_order(self, tmp_path):
        """多页按序拼接。"""
        f = tmp_path / "multi.pptx"
        _make_pptx(f, [
            {"body": "第一页内容"},
            {"body": "第二页内容"},
        ])
        text = await PptxParser().extract_text(f)
        assert text.index("第一页内容") < text.index("第二页内容")

    def test_corrupt_pptx_raises_value_error(self, tmp_path):
        """损坏 pptx → ValueError（ingest 沿既有路径转 A3001）。"""
        f = tmp_path / "corrupt.pptx"
        f.write_bytes(b"not a real pptx")
        with pytest.raises(ValueError, match="损坏|加密"):
            PptxParser()._extract(f)
