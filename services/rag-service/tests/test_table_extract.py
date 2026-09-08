"""内嵌表格抽取与结构化预览测试（V5，design D5；tasks 11.4）。

覆盖：表格行列结构保留 / 任意单元格命中（serialized_lines）/ 表格预览元数据返回
（kb_doc_assets content_json，gr.Dataframe 渲染单一数据源）/ 开关关闭走 V4 文本化（零资产）。
"""
import json
import uuid
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from app.parsers.asset_extractor import AssetExtractor
from app.services.asset_service import AssetService
from app.settings import settings


class TestTableAssetBuild:
    def test_structure_preserved(self):
        """行列结构保留：headers / data_rows / serialized_lines。"""
        asset = AssetExtractor._build_table_asset(3, [
            ["材料", "要求"],
            ["身份证", "原件"],
            ["合同", "复印件"],
        ])
        assert asset["type"] == "table"
        assert asset["page"] == 3
        assert asset["headers"] == ["材料", "要求"]
        assert asset["rows"] == [["身份证", "原件"], ["合同", "复印件"]]
        assert asset["serialized_lines"] == [
            "表格(第3页)[材料, 要求]: 材料:身份证, 要求:原件",
            "表格(第3页)[材料, 要求]: 材料:合同, 要求:复印件",
        ]
        assert asset.get("summary_line") == "表格(第3页)[材料, 要求]: 材料: 身份证, 合同"

    def test_location_prefix_and_column_summary(self):
        """行 content 含 [定位] 前缀；标识列汇总 chunk。"""
        asset = AssetExtractor._build_table_asset(1, [
            ["子系统", "技术类别", "说明"],
            ["RAG", "多模态解析", "a"],
            ["RAG", "语义分块", "b"],
            ["LLM", "向量化", "c"],
        ], location="3.1 RAG多模态知识库系统")
        assert asset["serialized_lines"][0].startswith(
            "表格(第1页)[3.1 RAG多模态知识库系统]:"
        )
        assert "技术类别:多模态解析" in asset["serialized_lines"][0]
        assert asset["summary_line"] is not None
        assert "技术类别: 多模态解析, 语义分块, 向量化" in asset["summary_line"]

    def test_any_cell_hit(self):
        """任意单元格命中：serialized_lines 含每个单元格值（行列不打散）。"""
        asset = AssetExtractor._build_table_asset(1, [
            ["a", "b"],
            ["v1", "v2"],
        ])
        joined = "\n".join(asset["serialized_lines"])
        assert "a:v1" in joined and "b:v2" in joined

    def test_empty_value_not_serialized(self):
        """空单元格不序列化（对应行可能被跳过）。"""
        asset = AssetExtractor._build_table_asset(1, [
            ["a", "b"],
            ["v1", ""],
        ])
        assert asset["serialized_lines"] == ["表格(第1页)[a, b]: a:v1"]


class TestTableAssetIngest:
    def _svc(self):
        """注入 AsyncMock 化的 batch_create（服务内 await，普通 MagicMock 会 TypeError）。"""
        chunk_repo = MagicMock()
        chunk_repo.batch_create = AsyncMock()
        chunk_repo.next_chunk_index = AsyncMock(return_value=1)
        asset_repo = MagicMock()
        asset_repo.batch_create = AsyncMock()
        return AssetService(
            doc_repo=MagicMock(), chunk_repo=chunk_repo, asset_repo=asset_repo,
            milvus_client=MagicMock(), es_client=MagicMock(), embed_client=MagicMock(),
        )

    @pytest.mark.asyncio
    async def test_table_chunk_and_preview_metadata(self, monkeypatch):
        """表格：序列化行以 chunk 入库 + kb_doc_assets 落 content_json（预览渲染源）。"""
        monkeypatch.setattr(settings, "KB_ASSETS_DIR", "/tmp")
        monkeypatch.setattr(settings, "KB_TABLE_EXTRACT_ENABLED", True)
        svc = self._svc()
        svc.embed.batch_embed = AsyncMock(return_value=[[0.1] * 4])
        svc.milvus.insert = AsyncMock(return_value=["vid"])
        svc.es.index_docs = AsyncMock()
        doc_id = uuid.uuid4()
        await svc.process_assets(doc_id, [{
            "type": "table", "page": 2,
            "headers": ["材料", "要求"],
            "rows": [["身份证", "原件"]],
            "serialized_lines": ["表格(第2页)[材料, 要求]: 材料:身份证, 要求:原件"],
        }], "regulation")
        # chunk 入库（保留行列结构不打散）
        chunk_args = svc.chunk_repo.batch_create.await_args.args[0]
        assert chunk_args[0]["content"] == "表格(第2页)[材料, 要求]: 材料:身份证, 要求:原件"
        svc.embed.batch_embed.assert_awaited_once()
        # kb_doc_assets 元数据：content_json 为表格行列结构化（gr.Dataframe 渲染源）
        assets = svc.asset_repo.batch_create.await_args.args[0]
        assert len(assets) == 1
        assert assets[0].asset_type == "table"
        assert assets[0].page == 2
        content = json.loads(assets[0].content_json)
        assert content["headers"] == ["材料", "要求"]
        assert content["rows"] == [["身份证", "原件"]]
        assert assets[0].storage_path.startswith(f"{doc_id}/")

    @pytest.mark.asyncio
    async def test_empty_table_not_ingested(self, monkeypatch):
        """空表/无序列化行 → 不落库（直接 return）。"""
        svc = self._svc()
        svc.embed.batch_embed = AsyncMock()
        await svc.process_assets(uuid.uuid4(), [{
            "type": "table", "page": 1,
            "headers": ["a"], "rows": [], "serialized_lines": [],
        }], "regulation")
        svc.embed.batch_embed.assert_not_awaited()
        svc.asset_repo.batch_create.assert_not_awaited()

    @pytest.mark.asyncio
    async def test_table_with_summary_chunk_batch(self, monkeypatch):
        """含技术类别列 → 额外写入列汇总 chunk。"""
        monkeypatch.setattr(settings, "KB_ASSETS_DIR", "/tmp")
        monkeypatch.setattr(settings, "KB_TABLE_EXTRACT_ENABLED", True)
        svc = self._svc()
        lines = [
            "表格(第1页)[章节]: 技术类别:多模态解析, 说明:a",
            "表格(第1页)[章节]: 技术类别:语义分块, 说明:b",
        ]
        summary = "表格(第1页)[章节]: 技术类别: 多模态解析, 语义分块"
        svc.embed.batch_embed = AsyncMock(return_value=[[0.1] * 4 for _ in range(3)])
        svc.milvus.insert = AsyncMock(return_value=["v1", "v2", "v3"])
        svc.es.index_docs = AsyncMock()
        doc_id = uuid.uuid4()
        await svc.process_assets(doc_id, [{
            "type": "table", "page": 1,
            "headers": ["技术类别", "说明"],
            "rows": [["多模态解析", "a"], ["语义分块", "b"]],
            "serialized_lines": lines,
            "summary_line": summary,
        }], "regulation")
        embed_inputs = svc.embed.batch_embed.await_args.args[0]
        assert len(embed_inputs) == 3
        assert summary in embed_inputs

    @pytest.mark.asyncio
    async def test_multi_row_table_single_batch_embed(self, monkeypatch):
        """多行表格一次 batch_embed，仍每行一个 chunk。"""
        monkeypatch.setattr(settings, "KB_ASSETS_DIR", "/tmp")
        monkeypatch.setattr(settings, "KB_TABLE_EXTRACT_ENABLED", True)
        svc = self._svc()
        lines = [
            "表格(第1页): a:1, b:2",
            "表格(第1页): a:3, b:4",
            "表格(第1页): a:5, b:6",
        ]
        svc.embed.batch_embed = AsyncMock(return_value=[[0.1] * 4 for _ in lines])
        svc.milvus.insert = AsyncMock(return_value=["v1", "v2", "v3"])
        svc.es.index_docs = AsyncMock()
        doc_id = uuid.uuid4()
        await svc.process_assets(doc_id, [{
            "type": "table", "page": 1,
            "headers": ["a", "b"],
            "rows": [["1", "2"], ["3", "4"], ["5", "6"]],
            "serialized_lines": lines,
        }], "regulation")
        svc.embed.batch_embed.assert_awaited_once()
        assert len(svc.embed.batch_embed.await_args.args[0]) == 3
        chunk_args = svc.chunk_repo.batch_create.await_args.args[0]
        assert len(chunk_args) == 3

    @pytest.mark.asyncio
    async def test_switches_off_v4_text_only(self, tmp_path, monkeypatch):
        """开关关闭：extract_with_assets 不执行资产抽取（V4 文本化，零资产）。"""
        monkeypatch.setattr(settings, "KB_IMAGE_EXTRACT_ENABLED", False)
        monkeypatch.setattr(settings, "KB_TABLE_EXTRACT_ENABLED", False)
        from pptx import Presentation
        from pptx.util import Inches

        from app.parsers.pptx_parser import PptxParser

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "合规制度"
        tb = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
        tb.text = "正文内容"
        f = tmp_path / "a.pptx"
        prs.save(str(f))

        with patch("app.parsers.asset_extractor.AssetExtractor") as mae:
            text, assets = await PptxParser().extract_with_assets(f)
        assert assets == []
        mae.extract.assert_not_called()
        assert "合规制度" in text
